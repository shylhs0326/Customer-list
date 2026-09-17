"""Excel input/output boundary. Source files are never modified."""
from datetime import datetime, date
from io import BytesIO
import re
import zipfile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from core import FIELDS, OUTPUT_FIELDS


def open_book(data, data_only=True):
    if len(data) > 30 * 1024 * 1024:
        raise ValueError('파일당 30MB 이하만 지원합니다.')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 250 * 1024 * 1024:
                raise ValueError('압축 해제 크기가 너무 큽니다. 필요한 시트만 저장해 주세요.')
        return load_workbook(BytesIO(data), read_only=True, data_only=data_only, keep_links=False)
    except (zipfile.BadZipFile, KeyError) as error:
        raise ValueError('일반 .xlsx 또는 .xlsm 파일을 선택해 주세요. 암호 파일은 지원하지 않습니다.') from error


def cell_text(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        digits = str(int(value))
        mask = (cell.number_format or '').split(';')[0]
        if re.fullmatch(r'[0\-() ]+', mask) and '0' in mask and value >= 0:
            digits = digits.zfill(mask.count('0'))
            if len(digits) == mask.count('0'):
                iterator = iter(digits)
                return ''.join(next(iterator) if c == '0' else c for c in mask)
        return digits
    return str(value).strip()


def inspect_book(data):
    book = open_book(data)
    try:
        sheets = []
        for sheet in book:
            if sheet.max_column and sheet.max_column > 300:
                raise ValueError('열이 300개를 넘습니다. 필요한 열만 새 파일에 저장해 주세요.')
            sheets.append(dict(name=sheet.title, rows=sheet.max_row,
                               preview=[[cell_text(c) for c in row] for row in sheet.iter_rows(max_row=min(sheet.max_row or 1, 30))]))
        return dict(sheets=sheets)
    finally:
        book.close()


def read_records(data, sheet_name, header_row, mapping, source):
    if mapping.get('customer_id') is None:
        raise ValueError('고객번호 열을 반드시 연결해 주세요.')
    if not isinstance(header_row, int) or not 1 <= header_row <= 30:
        raise ValueError('제목 행은 1~30 사이여야 합니다.')
    book = open_book(data)
    formulas = open_book(data, data_only=False)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError('시트를 찾을 수 없습니다.')
        sheet = book[sheet_name]
        indexes = [v for k, v in mapping.items() if k in FIELDS and v is not None]
        if any(not isinstance(v, int) or not 0 <= v < (sheet.max_column or 0) for v in indexes):
            raise ValueError('연결된 열 번호가 잘못되었습니다.')
        if len(indexes) != len(set(indexes)):
            raise ValueError('하나의 원본 열을 여러 항목에 중복 연결할 수 없습니다.')
        if sheet.max_row and sheet.max_row > 100001:
            raise ValueError('시트당 100,000행 이하만 지원합니다.')
        rows = []
        for number, (cells, original) in enumerate(zip(sheet.iter_rows(min_row=header_row + 1),
                                                       formulas[sheet_name].iter_rows(min_row=header_row + 1)), header_row + 1):
            if not any(c.value is not None for c in original):
                continue
            record = {key: '' for key in FIELDS}
            warnings = []
            for key, index in mapping.items():
                if key not in FIELDS or index is None:
                    continue
                cell = cells[index]
                if original[index].data_type == 'f' and cell.value is None:
                    warnings.append(f'{FIELDS[key]} 수식 결과 없음: Excel에서 계산 후 저장 필요')
                if cell.data_type == 'e':
                    warnings.append(f'{FIELDS[key]} Excel 오류 값')
                    continue
                record[key] = cell_text(cell)
                if key in ('customer_id', 'machine', 'phone', 'mobile') and isinstance(cell.value, (int, float)):
                    if len(str(int(cell.value))) > 15:
                        warnings.append(f'{FIELDS[key]} 숫자 정밀도 손실 가능: 원본 텍스트 확인')
                    if key in ('phone', 'mobile') and not record[key].startswith(('0', '+')):
                        warnings.append('숫자 연락처의 앞자리 0 유실 가능')
            record.update(source=f'{source} / {sheet_name} / {number}행', warnings=warnings)
            rows.append(record)
        if not rows:
            raise ValueError('선택한 제목 행 아래에 데이터가 없습니다.')
        return rows
    finally:
        book.close()
        formulas.close()


def add_table(book, name, columns, rows):
    sheet = book.create_sheet(name)
    sheet.append([label for key, label in columns])
    for row_index, row in enumerate(rows, 2):
        values = []
        for key, label in columns:
            value = row.get(key, '')
            if isinstance(value, list):
                value = ' / '.join(str(v) for v in value)
            values.append(str(value) if value is not None else '')
        sheet.append(values)
        for column_index in range(1, len(columns) + 1):
            cell = sheet.cell(row_index, column_index)
            cell.data_type = 's'  # Never execute input as an Excel formula.
            cell.number_format = '@'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='183E45')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for index, (key, label) in enumerate(columns, 1):
        width = 34 if key in ('address', 'email', 'source', 'all_sources', 'selection', 'reasons', 'warnings') else 22
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cells in sheet.iter_rows(min_row=2):
        max_lines = 1
        for cell in cells:
            cell.font = Font(name='맑은 고딕', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F6F5')
            max_lines = max(max_lines, len(str(cell.value or '').splitlines()))
        sheet.row_dimensions[cells[0].row].height = min(150, max(32, max_lines * 16))
    return sheet


def export_result(result):
    book = Workbook()
    book.remove(book.active)
    cols = [(key, '성명' if key == 'name' else FIELDS[key]) for key in OUTPUT_FIELDS]
    extras = [('address', '현재 주소'), ('priority', '선정 순위'), ('position', '직급'), ('last_name', '성'), ('first_name', '이름'), ('mobile', '휴대폰'), ('updated', '선정 기준일 (기입날짜/Install Date)'), ('selection', '선정 근거'),
              ('confirmed_at', '사용자 확인일'), ('source', '담당자 원본 출처'), ('all_sources', '통합 원본 출처')]
    add_table(book, '설문 대상', cols + extras, [c for c in result['customers'] if c['status'] == '설문 대상'])
    reviews = [c for c in result['customers'] if c['status'] != '설문 대상']
    reviews += [dict(r, reasons=[r['reason']]) for r in result['rejected']]
    add_table(book, '확인 필요', cols + [('reasons', '확인 사유')] + extras, reviews)
    add_table(book, '설치 기계', [(k, FIELDS[k]) for k in ('customer_id', 'company', 'machine', 'model')] + [('source', '출처')], result['machines'])
    add_table(book, '원본 통합', [(k, '성명' if k == 'name' else v) for k, v in FIELDS.items()] + [('source', '출처'), ('warnings', '읽기 주의사항')], result['records'])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def sample_book(which):
    book = Workbook()
    sheet = book.active
    sheet.title = '고객정보'
    sheet.append(list(FIELDS.values()))
    sheet['E1'] = '성명'
    sheet['K1'] = '기입날짜' if which == 1 else 'Install Date'
    # Entirely fictional data. example.com is reserved for examples.
    rows = [
        ['M001', '모델 A', '000101', '가상테크', '김총무', '010-0000-1001', 'kim@example.com', '서울', '총무팀', '', '2026-08-01'],
        ['M002', '모델 B', '000101', '가상테크', '김총무', '010-0000-1001', 'kim@example.com', '서울', '총무팀', '', '2026-08-01'],
        ['M003', '모델 A', '000102', '가상산업', '이인사', '010-0000-1002', '', '부산', '인사팀', '', '2026-08-05'],
        ['M004', '모델 C', '000103', '가상상사', '박지원', '010-0000-1003', 'park@example.com', '대전', '경영지원팀', '', ''],
    ] if which == 1 else [
        ['M001', '모델 A', '000101', '가상테크', '최구매', '010-0000-2001', 'choi@example.com', '서울', '구매팀', '예', '2026-09-01'],
        ['M003', '모델 A', '000102', '가상산업', '이인사', '010-0000-1002', 'lee@example.com', '부산', '인사팀', '', '2026-09-02'],
        ['M005', '모델 D', '000104', '가상물류', '정개발', '010-0000-1004', 'jung@example.com', '인천', '개발팀', '아니오', '2026-09-02'],
        ['M006', '모델 A', '', '고객번호누락예제', '서총무', '010-0000-1005', 'seo@example.com', '서울', '총무팀', '', '2026-09-02'],
    ]
    for row in rows:
        row[8] = row[7] + ' 예제빌딩 3층 ' + row[8]
        row[9] = '과장'
        row.extend([row[4][:1], row[4][1:]])
        row.append(row[5])
        row[5] = row[5].replace('010-', '02-', 1)
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()
