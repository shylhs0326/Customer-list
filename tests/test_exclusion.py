import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from core import build_result
from excel_io import export_result, read_exclusion_ids


def record(customer_id='C-001'):
    return {
        'machine': 'M-001', 'model': 'Model-A', 'customer_id': customer_id,
        'company': '가상회사', 'name': '김총무', 'phone': '02-0000-0001',
        'email': 'kim@example.com', 'region': '서울',
        'address': '서울 구매팀', 'position': '과장',
        'updated': '2026-09-01', 'last_name': '김', 'first_name': '총무',
        'mobile': '010-0000-0001', 'source': '자료1', 'warnings': [],
    }


class ExclusionTests(unittest.TestCase):
    def test_exclusion_workbook_reads_customer_number_column(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['고객번호'])
        sheet.append(['000101'])
        sheet.append(['000102'])
        output = BytesIO()
        book.save(output)

        self.assertEqual(read_exclusion_ids(output.getvalue()), ['000101', '000102'])

    def test_headerless_exclusion_workbook_uses_first_column(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['000201'])
        sheet.append(['000202'])
        output = BytesIO()
        book.save(output)

        self.assertEqual(read_exclusion_ids(output.getvalue()), ['000201', '000202'])

    def test_customer_id_in_exclusion_list_is_marked_excluded(self):
        result = build_result([record()], excluded_ids={'C-001'})

        customer = result['customers'][0]
        self.assertEqual(customer['status'], '설문 제외')
        self.assertIn('제외 목록에 포함', customer['reasons'])

        saved = load_workbook(BytesIO(export_result(result)), read_only=True)
        self.assertIn('설문 제외', saved.sheetnames)
        self.assertEqual(saved['설문 제외'].max_row, 2)

    def test_customer_id_not_in_exclusion_list_keeps_normal_selection(self):
        result = build_result([record()], excluded_ids={'C-999'})

        self.assertEqual(result['customers'][0]['status'], '설문 대상')


if __name__ == '__main__':
    unittest.main()
