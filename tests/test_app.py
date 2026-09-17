from io import BytesIO
import json
import threading
import unittest
import urllib.request
import urllib.error
from openpyxl import load_workbook
from app import SurveyServer
from core import FIELDS


class ApplicationFlow(unittest.TestCase):
    def setUp(self):
        self.server = SurveyServer(('127.0.0.1', 0))
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f'http://127.0.0.1:{self.server.server_port}'

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()

    def post(self, endpoint, payload, token=True):
        headers = {'Content-Type': 'application/json'}
        if token:
            headers['X-Session-Token'] = self.server.token
        request = urllib.request.Request(self.url + '/api/' + endpoint,
                                         data=json.dumps(payload).encode(), headers=headers)
        with urllib.request.urlopen(request) as response:
            content = response.read()
            return json.loads(content) if response.headers['Content-Type'].startswith('application/json') else content

    def build(self):
        self.post('demo', {})
        mapping = {key: index for index, key in enumerate(FIELDS)}
        return self.post('build', {'inputs': [dict(slot=s, sheet='고객정보', header=1, mapping=mapping)
                                             for s in ('1', '2')]})

    def test_demo_build_review_export_roundtrip(self):
        result = self.build()
        self.assertEqual(len(result['customers']), 4)
        self.assertEqual(len(result['rejected']), 1)
        self.assertEqual(sum(c['status'] == '설문 대상' for c in result['customers']), 2)
        customer = next(c for c in result['customers'] if c['customer_id'] == '000103')
        result = self.post('approve', dict(customer_id='000103', values=customer, confirmed=True))
        book = load_workbook(BytesIO(self.post('export', {})), data_only=False)
        self.assertEqual(book['설문 대상'].max_row, 4)
        self.assertEqual(book['확인 필요'].max_row, 3)
        self.assertEqual(book['설치 기계'].max_row, 6)
        self.assertEqual(book['원본 통합'].max_row, 9)
        self.assertEqual(book['설문 대상']['C2'].value, '000101')
        self.assertEqual(book['설문 대상']['A2'].value, 'M001\nM002')

    def test_confirmation_is_required_and_errors_do_not_mutate(self):
        result = self.build()
        customer = result['customers'][2]
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.post('approve', dict(customer_id=customer['customer_id'], values=customer, confirmed=False))
        self.assertEqual(error.exception.code, 400)
        self.assertEqual(self.server.result['customers'][2]['status'], '확인 필요')

    def test_other_origin_cannot_submit(self):
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.post('demo', {}, token=False)
        self.assertEqual(error.exception.code, 403)

    def test_new_file_selection_invalidates_export(self):
        self.build()
        self.post('demo', {})
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.post('export', {})
        self.assertEqual(error.exception.code, 400)


if __name__ == '__main__':
    unittest.main()
