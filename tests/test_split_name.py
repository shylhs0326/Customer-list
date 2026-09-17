import unittest
from core import build_result, approve_customer
from test_address import contact
from excel_io import export_result
from openpyxl import load_workbook
from io import BytesIO


class SplitNames(unittest.TestCase):
    def test_parts_form_full_name(self):
        c = build_result([contact(name='', last_name='남궁', first_name='민수')])['customers'][0]
        self.assertEqual(c['name'], '남궁민수')
        self.assertEqual(c['last_name'], '남궁')
        self.assertEqual(c['status'], '설문 대상')

    def test_missing_part_requires_review(self):
        for last, first in [('김',''), ('','민수')]:
            c = build_result([contact(name='', last_name=last, first_name=first)])['customers'][0]
            self.assertEqual(c['status'], '확인 필요')

    def test_different_surnames_do_not_collapse(self):
        c = build_result([contact(name='', last_name='김', first_name='민수'),
                          contact(name='', last_name='이', first_name='민수')])['customers'][0]
        self.assertEqual(c['status'], '확인 필요')

    def test_manual_update_rebuilds_full_name(self):
        result = build_result([contact(name='', last_name='김', first_name='민수')])
        approve_customer(result, '001', contact(last_name='박', first_name='지원'))
        self.assertEqual(result['customers'][0]['name'], '박지원')

    def test_export_has_separate_name_columns(self):
        result = build_result([contact(name='', last_name='남궁', first_name='민수')])
        sheet = load_workbook(BytesIO(export_result(result)))['설문 대상']
        headers = [c.value for c in sheet[1]]
        for label, expected in [('성','남궁'), ('이름','민수'), ('성명','남궁민수')]:
            self.assertEqual(sheet.cell(2, headers.index(label)+1).value, expected)

if __name__ == '__main__':
    unittest.main()
