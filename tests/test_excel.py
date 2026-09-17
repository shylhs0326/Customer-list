import unittest
from io import BytesIO
from openpyxl import Workbook, load_workbook
from excel_io import inspect_book, read_records, export_result
from core import build_result


def fixture():
    book = Workbook()
    sheet = book.active
    sheet.title = '고객'
    sheet.append(['고객 자료'])
    sheet.append(['고객코드', '성명', '전화', '메일', '부서명', '수정일'])
    sheet.append([12, '=홍길동', '01012345678', 'a@example.com', '총무팀', '2026-09-01'])
    sheet['A3'].number_format = '000000'
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class ExcelBoundary(unittest.TestCase):
    def test_read_column_mapping_and_formatted_identifier(self):
        rows = read_records(fixture(), '고객', 2, {'customer_id': 0, 'phone': 2}, 'A.xlsx')
        self.assertEqual(rows[0]['customer_id'], '000012')
        self.assertEqual(rows[0]['phone'], '01012345678')
        self.assertIn('3행', rows[0]['source'])

    def test_inspection_has_sheet_and_preview(self):
        info = inspect_book(fixture())
        self.assertEqual(info['sheets'][0]['name'], '고객')
        self.assertEqual(info['sheets'][0]['preview'][1][0], '고객코드')

    def test_formula_without_cached_value_does_not_look_complete(self):
        rows = read_records(fixture(), '고객', 2, {'customer_id': 0, 'name': 1}, 'A.xlsx')
        self.assertEqual(rows[0]['name'], '')
        self.assertTrue(rows[0]['warnings'])

    def test_export_untrusted_text_stays_text(self):
        result = build_result([dict(customer_id='001', company='=1+1', name='담당')])
        output = export_result(result)
        book = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(book.sheetnames, ['설문 대상', '확인 필요', '설치 기계', '원본 통합'])
        self.assertEqual(book['확인 필요']['D2'].value, '=1+1')
        self.assertEqual(book['확인 필요']['D2'].data_type, 's')

    def test_customer_mapping_is_required(self):
        with self.assertRaises(ValueError):
            read_records(fixture(), '고객', 2, {'name': 1}, 'A.xlsx')


if __name__ == '__main__':
    unittest.main()
