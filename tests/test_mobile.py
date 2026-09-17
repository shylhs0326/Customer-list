import unittest
from io import BytesIO
from openpyxl import load_workbook
from core import build_result, approve_customer
from excel_io import export_result
from test_address import contact


class SeparateNumbers(unittest.TestCase):
    def test_mobile_only_is_contactable(self):
        c = build_result([contact(phone='', mobile='01012345678')])['customers'][0]
        self.assertEqual(c['status'], '설문 대상')
        self.assertEqual(c['mobile'], '01012345678')
        self.assertEqual(c['phone'], '')

    def test_secondary_candidate_can_have_only_mobile(self):
        c = build_result([contact(address='서울', position='대표', phone='', mobile='01012345678')])['customers'][0]
        self.assertEqual(c['priority'], 2)
        self.assertEqual(c['status'], '설문 대상')

    def test_both_missing_requires_review(self):
        self.assertEqual(build_result([contact(phone='', mobile='')])['customers'][0]['status'], '확인 필요')

    def test_invalid_mobile_requires_review(self):
        self.assertEqual(build_result([contact(mobile='bad')])['customers'][0]['status'], '확인 필요')

    def test_mobile_conflicts_require_review(self):
        c = build_result([contact(mobile='01011112222'), contact(mobile='01033334444')])['customers'][0]
        self.assertEqual(c['status'], '확인 필요')

    def test_edit_and_export_keep_numbers_separate(self):
        result = build_result([contact(phone='0212345678', mobile='01012345678')])
        approve_customer(result, '001', contact(phone='0233334444', mobile='01055556666'))
        sheet = load_workbook(BytesIO(export_result(result)))['설문 대상']
        headers = [c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index('전화')+1).value, '0233334444')
        self.assertEqual(sheet.cell(2, headers.index('휴대폰')+1).value, '01055556666')

if __name__ == '__main__':
    unittest.main()
