import unittest
from io import BytesIO
from openpyxl import load_workbook
from core import build_result, approve_customer
from excel_io import export_result, read_records, sample_book
from test_address import contact


class PositionField(unittest.TestCase):
    def test_job_title_is_preserved_without_yes_no_validation(self):
        customer = build_result([contact(position='과장')])['customers'][0]
        self.assertEqual(customer.get('position'), '과장')
        self.assertEqual(customer['status'], '설문 대상')

    def test_job_title_does_not_imply_purchase_eligibility(self):
        customer = build_result([contact(address='서울 영업팀', position='부장')])['customers'][0]
        self.assertEqual(customer['status'], '확인 필요')

    def test_manual_edit_and_export_keep_job_title(self):
        result = build_result([contact(position='사원')])
        approve_customer(result, '001', contact(position='대리'))
        self.assertEqual(result['customers'][0].get('position'), '대리')
        book = load_workbook(BytesIO(export_result(result)))
        headers = [c.value for c in book['설문 대상'][1]]
        self.assertIn('직급', headers)
        self.assertNotIn('구매권한', headers)
        self.assertEqual(book['설문 대상'].cell(2, headers.index('직급')+1).value, '대리')

    def test_sample_can_map_job_title(self):
        data = sample_book(1)
        book = load_workbook(BytesIO(data))
        self.assertEqual(book.active['J1'].value, '직급')
        rows = read_records(data, '고객정보', 1, {'customer_id': 2, 'position': 9}, 'sample.xlsx')
        self.assertTrue(rows[0].get('position'))

if __name__ == '__main__':
    unittest.main()
