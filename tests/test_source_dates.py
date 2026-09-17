import unittest
from io import BytesIO
from openpyxl import load_workbook
from core import build_result, FIELDS
from excel_io import sample_book, read_records


class SourceDates(unittest.TestCase):
    def test_sample_headers_follow_source(self):
        for slot, label in [(1, '기입날짜'), (2, 'Install Date')]:
            book = load_workbook(BytesIO(sample_book(slot)))
            self.assertEqual(book.active['K1'].value, label)

    def test_newest_wins_in_either_source(self):
        for date1, date2, expected in [('2026-09-10','2026-09-01','자료1담당'),
                                        ('2026-09-01','2026-09-10','자료2담당')]:
            records = []
            for slot, date_value in [(1,date1),(2,date2)]:
                book = load_workbook(BytesIO(sample_book(slot)))
                sheet = book.active
                sheet.delete_rows(3, sheet.max_row)
                sheet['E2'] = f'자료{slot}담당'
                sheet['L2'] = None
                sheet['M2'] = None
                sheet['K2'] = date_value
                stream = BytesIO()
                book.save(stream)
                records += read_records(stream.getvalue(), '고객정보', 1,
                                        {key:i for i,key in enumerate(FIELDS)}, f'자료{slot}.xlsx')
            result = build_result(records)['customers'][0]
            self.assertEqual(result['name'], expected)
            self.assertEqual(result['updated'], '2026-09-10')
            self.assertEqual(result['status'], '설문 대상')

if __name__ == '__main__':
    unittest.main()
